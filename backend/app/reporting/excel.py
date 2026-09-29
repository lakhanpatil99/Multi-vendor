"""Excel report generation (openpyxl). Structured multi-sheet workbook."""
from __future__ import annotations

from typing import Any


def excel_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:  # pragma: no cover
        return False


def build_excel(title: str, preview: dict[str, Any], findings: list[dict],
                remediations: list[dict], audit: list[dict]) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="0F2740")
    header_font = Font(color="FFFFFF", bold=True)

    def _sheet(name: str, headers: list[str], rows: list[list]):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            ws.append(row)
        return ws

    # Summary
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Report", title])
    summary.append(["Compliance Score", f"{preview.get('compliance_score', 0)}%"])
    summary.append(["Findings", len(findings)])
    summary.append(["Executive Summary", preview.get("executive_summary", "")])

    _sheet("Framework Mapping",
           ["Framework", "Score", "Pass", "Fail"],
           [[f["framework"], f["score"], f["passed"], f["failed"]]
            for f in preview.get("framework_results", [])])

    _sheet("Findings",
           ["ID", "Title", "Severity", "Status", "Category", "Actual", "Expected"],
           [[f.get("id"), f.get("title"), f.get("severity"), f.get("status"),
             f.get("category"), f.get("actual_value"), f.get("expected_value")]
            for f in findings])

    _sheet("Remediation",
           ["Finding", "Vendor", "Severity", "Status", "Commands"],
           [[r.get("title"), r.get("vendor"), r.get("severity"), r.get("status"),
             " ; ".join(r.get("commands", []))] for r in remediations])

    _sheet("Audit",
           ["Event", "Title", "Actor", "Timestamp"],
           [[a.get("event_type"), a.get("title"), a.get("actor"), a.get("created_at")]
            for a in audit])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
